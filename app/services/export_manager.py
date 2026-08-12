"""
This module contains the ExportManager class for generating PDF, Word, and JSON
exports of conversations.
"""

import json
import os
import platform
from io import BytesIO

# ReportLab - PDF generation
try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# Python-docx for Word document generation
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class ExportManager:
    def __init__(self, db_service, export_config):
        """Initializes the export manager with injected database service and export configuration."""
        self.db = db_service
        self.config = export_config

    def _register_modern_fonts(self):
        """Registers system TTF fonts for ReportLab PDF generation."""
        try:
            if platform.system() == "Windows":
                font_paths = [
                    "C:/Windows/Fonts/segoeui.ttf",
                    "C:/Windows/Fonts/calibri.ttf",
                    "C:/Windows/Fonts/arial.ttf"
                ]
                bold_paths = [
                    "C:/Windows/Fonts/segoeuib.ttf",
                    "C:/Windows/Fonts/calibrib.ttf",
                    "C:/Windows/Fonts/arialbd.ttf"
                ]
            elif platform.system() == "Darwin":
                font_paths = [
                    "/System/Library/Fonts/SF-Pro-Text-Regular.otf",
                    "/System/Library/Fonts/Helvetica.ttc",
                    "/Library/Fonts/Arial.ttf"
                ]
                bold_paths = [
                    "/System/Library/Fonts/SF-Pro-Text-Bold.otf",
                    "/System/Library/Fonts/Helvetica-Bold.ttc",
                    "/Library/Fonts/Arial Bold.ttf"
                ]
            else:
                # Standard Linux font paths
                font_paths = [
                    "/usr/share/fonts/truetype/ubuntu/Ubuntu-R.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
                ]
                bold_paths = [
                    "/usr/share/fonts/truetype/ubuntu/Ubuntu-B.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
                ]

            for regular, bold in zip(font_paths, bold_paths):
                try:
                    if os.path.exists(regular) and os.path.exists(bold):
                        pdfmetrics.registerFont(TTFont('ModernFont', regular))
                        pdfmetrics.registerFont(TTFont('ModernFont-Bold', bold))
                        return 'ModernFont'
                except Exception:
                    continue
            return self.config.fallback_font
        except Exception:
            return self.config.fallback_font

    def create_json(self, conversation_id: int, session_id: str, summary: str = None) -> BytesIO:
        """Exports conversation history as a JSON stream."""
        conv_info = self.db.get_conversation_info(conversation_id, session_id)
        if not conv_info:
            return None

        character, title = conv_info
        messages = self.db.get_messages(conversation_id, session_id)

        data = {
            "document": self.config.document_title,
            "character": character,
            "title": title,
            "summary": summary,
            "messages": [{"question": q, "answer": a} for q, a in messages]
        }

        return BytesIO(json.dumps(data, indent=4, ensure_ascii=False).encode("utf-8"))

    def create_pdf(self, conversation_id: int, session_id: str, summary: str = None) -> BytesIO:
        """Generates a formatted PDF document of the conversation."""
        if not REPORTLAB_AVAILABLE:
            raise ImportError("ReportLab is required for PDF generation: pip install reportlab")

        conv_info = self.db.get_conversation_info(conversation_id, session_id)
        if not conv_info:
            return None

        character, title = conv_info
        messages = self.db.get_messages(conversation_id, session_id)

        modern_font = self._register_modern_fonts()
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=50, leftMargin=50,
                                topMargin=50, bottomMargin=50)

        styles = getSampleStyleSheet()
        font_bold = f'{modern_font}-Bold' if modern_font != self.config.fallback_font else f'{self.config.fallback_font}-Bold'

        title_style = ParagraphStyle(
            'ModernTitle',
            parent=styles['Heading1'],
            fontSize=24,
            alignment=1,
            spaceAfter=30,
            textColor=HexColor(self.config.pdf_title_color),
            fontName=font_bold
        )
        subtitle_style = ParagraphStyle(
            'ModernSubtitle',
            parent=styles['Heading2'],
            fontSize=16,
            alignment=1,
            spaceAfter=20,
            textColor=HexColor(self.config.pdf_subtitle_color),
            fontName=font_bold
        )
        question_style = ParagraphStyle(
            'ModernQuestion',
            parent=styles['Normal'],
            fontSize=12,
            spaceAfter=8,
            spaceBefore=15,
            textColor=HexColor(self.config.pdf_question_color),
            fontName=font_bold
        )
        answer_style = ParagraphStyle(
            'ModernAnswer',
            parent=styles['Normal'],
            fontSize=11,
            spaceAfter=15,
            leading=16,
            textColor=HexColor(self.config.pdf_answer_color),
            fontName=modern_font
        )

        content = []
        content.append(Paragraph(f"🧙‍♂ {self.config.document_title}", title_style))
        content.append(Paragraph(f"Character: {character}", subtitle_style))
        content.append(Spacer(1, 30))

        if summary:
            summary_clean = str(summary).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            content.append(Paragraph("<b>📋 Conversation Summary:</b>", subtitle_style))
            content.append(Paragraph(summary_clean, answer_style))
            content.append(Spacer(1, 20))

        for i, (question, answer) in enumerate(messages, 1):
            def clean_text(text):
                return str(text).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

            content.append(Paragraph(f"<b>❓ Question {i}:</b> {clean_text(question)}", question_style))
            content.append(Paragraph(f"<b>💬 {character}:</b> {clean_text(answer)}", answer_style))
            
            if i < len(messages):
                content.append(Spacer(1, 10))

        try:
            doc.build(content)
            buffer.seek(0)
            return buffer
        except Exception as e:
            raise Exception(f"Failed to generate PDF: {str(e)}")

    def create_word(self, conversation_id: int, session_id: str, summary: str = None) -> BytesIO:
        """Generates a formatted Microsoft Word (.docx) document of the conversation."""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for Word generation: pip install python-docx")

        conv_info = self.db.get_conversation_info(conversation_id, session_id)
        if not conv_info:
            return None

        character, title = conv_info
        messages = self.db.get_messages(conversation_id, session_id)

        doc = Document()

        title_para = doc.add_heading(f'🧙‍♂ {self.config.document_title}', 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        char_para = doc.add_heading(f'Character: {character}', level=1)
        char_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph('')

        if summary:
            doc.add_heading('📋 Conversation Summary', level=2)
            doc.add_paragraph(summary)
            doc.add_paragraph('')

        for i, (question, answer) in enumerate(messages, 1):
            q_para = doc.add_paragraph()
            q_run = q_para.add_run(f'❓ Question {i}: ')
            q_run.bold = True
            q_run.font.color.rgb = RGBColor(*self.config.word_question_color)
            q_para.add_run(question)

            a_para = doc.add_paragraph()
            a_run = a_para.add_run(f'💬 {character}: ')
            a_run.bold = True
            a_run.font.color.rgb = RGBColor(*self.config.word_answer_color)
            a_para.add_run(answer)
            doc.add_paragraph('')

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer